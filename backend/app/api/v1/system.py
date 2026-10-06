"""System administration for the pilot (Phase 7): health, background jobs, message outbox, and bulk onboarding of
schools and users from CSV (always previewed first: nothing is written unless `commit=true`)."""
import csv
import io
import re
import uuid
from collections import Counter

from fastapi import APIRouter, Depends, File, Request, UploadFile
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.database import get_db
from app.core.deps import Principal, require, require_any
from app.core.errors import AppError
from app.core.security import hash_password, new_opaque_token
from app.core.timeutil import utcnow
from app.models import Organization, OrgType, OutboundMessage, Role, SystemSetting, User, UserRole, UserStatus
from app.schemas.common import normalize_phone
from app.seed.matrix import CONFLICTING_ROLE_PAIRS
from app.services import audit, notify

router = APIRouter(prefix="/system", tags=["system administration"])
MAX_CSV = 2 * 1024 * 1024


SHEETS = {"schools": ("Schools",), "users": ("Staff", "Staff accounts", "Users"), "prices": ("Prices", "Reference prices")}


async def _read_csv(file: UploadFile, kind: str = "") -> list[dict]:
    """CSV, or an Excel workbook (.xlsx): the sheet named after the import (e.g. "Schools") or else the first sheet."""
    raw = await file.read(MAX_CSV * 5 + 1)
    if raw[:2] == b"PK":
        return _read_xlsx(raw, kind)
    if len(raw) > MAX_CSV:
        raise AppError(413, "FILE_TOO_LARGE", "CSV files must be under 2 MB.")
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    reader = csv.DictReader(io.StringIO(text))
    return [{(k or "").strip().lower(): (v or "").strip() for k, v in row.items()} for row in reader]


def _read_xlsx(raw: bytes, kind: str) -> list[dict]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise AppError(422, "VALIDATION_ERROR", "This file could not be read as an Excel workbook (.xlsx).")
    names = {n.lower(): n for n in wb.sheetnames}
    sheet = next((wb[names[s.lower()]] for s in SHEETS.get(kind, ()) if s.lower() in names), wb.worksheets[0])
    rows = list(sheet.iter_rows(values_only=True))
    # the header is the first row that has the required first column (lets a workbook keep a title or notes above it)
    hi = next((i for i, r in enumerate(rows[:10]) if r and any(str(c or "").strip().lower() in
              ("county", "full_name", "code") for c in r)), 0)
    head = [str(c or "").strip().lower() for c in rows[hi]] if rows else []

    def txt(v):
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return str(v).strip()
    out = [{h: txt(v) for h, v in zip(head, r) if h} for r in rows[hi + 1:]]
    return [r for r in out if any(r.values())]


def _slug(s: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "-", s.upper()).strip("-")[:20]


# ---------------- status ----------------
@router.get("/status")
def status(p: Principal = Depends(require_any("iam:edit", "md:edit")), db: Session = Depends(get_db)):
    jobs = db.scalar(select(SystemSetting).where(SystemSetting.key == "jobs.last_run"))
    counts = dict(db.execute(select(OutboundMessage.status, func.count()).group_by(OutboundMessage.status)).all())
    failed = db.scalars(select(OutboundMessage).where(OutboundMessage.status == "failed").order_by(OutboundMessage.created_at.desc()).limit(20)).all()
    return {"env": settings.app_env, "database": db.bind.dialect.name, "notify_backend": settings.notify_backend,
            "email": "smtp" if settings.smtp_host else "console", "rate_limits": settings.rate_limit_enabled, "https_cookies": settings.cookie_secure,
            "jobs": (jobs.value if jobs else {}), "outbox": {"queued": counts.get("queued", 0), "sent": counts.get("sent", 0), "failed": counts.get("failed", 0)},
            "failed_messages": [{"id": m.id, "channel": m.channel, "to": notify.mask_phone(m.to) if m.channel == "sms" else m.to, "subject": m.subject,
                                 "attempts": m.attempts, "error": m.last_error, "created_at": m.created_at} for m in failed],
            "counts": {"users": db.scalar(select(func.count()).select_from(User)),
                       "schools": db.scalar(select(func.count()).select_from(Organization).where(Organization.type == OrgType.school))}}


@router.post("/outbox/{mid}/retry")
def retry(mid: uuid.UUID, p: Principal = Depends(require_any("iam:edit", "md:edit")), db: Session = Depends(get_db)):
    m = db.get(OutboundMessage, mid)
    if m is None or m.status != "failed":
        raise AppError(404, "NOT_FOUND", "No failed message with that id.")
    m.status, m.attempts, m.next_attempt_at = "queued", 0, utcnow()
    ok = notify.attempt(m)
    db.commit()
    return {"sent": ok, "error": m.last_error}


@router.post("/jobs/{task}/run")
def run_job(task: str, request: Request, p: Principal = Depends(require_any("iam:edit", "md:edit")), db: Session = Depends(get_db)):
    from app import jobs
    if task not in jobs.TASKS:
        raise AppError(404, "NOT_FOUND", "Unknown job.")
    audit.record(db, action="RUN_JOB", entity="job", entity_id=task, user=p.user, request=request)
    db.commit()
    return jobs.run([task])[task]


# ---------------- school import ----------------
@router.post("/import/schools")
async def import_schools(request: Request, commit: bool = False, file: UploadFile = File(...), p: Principal = Depends(require("md:create")),
                         db: Session = Depends(get_db)):
    """CSV columns: county, sub_county, school_code, school_name; optional: cluster, enrolment, lat, lng.
    Counties must already exist (by name or code). Sub-counties and clusters are created when missing; existing
    school codes are updated."""
    rows = await _read_csv(file, "schools")
    need = {"county", "sub_county", "school_code", "school_name"}
    if not rows or not need <= set(rows[0]):
        raise AppError(422, "VALIDATION_ERROR", "Columns needed: county, sub_county, school_code, school_name (optional: cluster, enrolment, lat, lng).")
    counties = {c.name.lower(): c for c in db.scalars(select(Organization).where(Organization.type == OrgType.county)).all()}
    counties |= {c.code.lower(): c for c in counties.values()}
    codes = Counter(r["school_code"].upper() for r in rows)
    result, created, updated = [], 0, 0
    for i, r in enumerate(rows, start=2):
        errs = []
        county = counties.get(r["county"].lower())
        if county is None:
            errs.append(f"Unknown county '{r['county']}'")
        code = r["school_code"].upper()
        if not code or not r["school_name"] or not r["sub_county"]:
            errs.append("School code, name and sub-county are required")
        if codes[code] > 1:
            errs.append("School code repeated in the file")
        enrol = None
        if r.get("enrolment"):
            try:
                enrol = int(float(r["enrolment"]))
                if enrol < 0 or enrol > 20000:
                    raise ValueError
            except ValueError:
                errs.append("Enrolment must be a number")
        gps = None
        if r.get("lat") or r.get("lng"):
            try:
                gps = {"lat": round(float(r["lat"]), 6), "lng": round(float(r["lng"]), 6), "source": "import",
                       "by": p.user.full_name, "at": utcnow().isoformat()}
                if not (-4.9 <= gps["lat"] <= 5.1 and 33.8 <= gps["lng"] <= 42.0):
                    raise ValueError
            except (ValueError, KeyError, TypeError):
                errs.append("lat/lng must be decimal degrees inside Kenya (e.g. -1.80, 37.62; check the minus sign and that they are not swapped)")
        existing = db.scalar(select(Organization).where(Organization.code == code)) if code else None
        if existing is not None and existing.type != OrgType.school:
            errs.append(f"Code {code} is already used by a {existing.type.value}")
        action = "error" if errs else ("update" if existing else "create")
        result.append({"line": i, "school_code": code, "school_name": r["school_name"], "county": county.name if county else r["county"],
                       "action": action, "errors": errs})
        if errs or not commit:
            continue

        def ensure(type_, name, parent):
            c = f"{parent.code}-{_slug(name)}"
            o = db.scalar(select(Organization).where(Organization.code == c))
            if o is None:
                o = Organization(type=type_, name=name, code=c, parent_id=parent.id, meta={}, created_by=p.id)
                db.add(o)
                db.flush()
            return o
        sc = ensure(OrgType.sub_county, r["sub_county"], county)
        parent = ensure(OrgType.school_cluster, r["cluster"], sc) if r.get("cluster") else sc
        meta = dict(existing.meta or {}) if existing else {}
        if enrol is not None:
            meta["enrolment"] = enrol
        if r.get("nemis_code"):
            meta["nemis_code"] = r["nemis_code"].strip().upper()[:30]
        if gps:
            meta["gps"] = gps
        if existing:
            existing.name, existing.parent_id, existing.meta = r["school_name"], parent.id, meta
            updated += 1
        else:
            db.add(Organization(type=OrgType.school, name=r["school_name"], code=code, parent_id=parent.id, meta=meta, created_by=p.id))
            created += 1
    ok = all(x["action"] != "error" for x in result)
    if commit:
        if not ok:
            db.rollback()
            raise AppError(422, "IMPORT_ERRORS", "Fix the errors in the file first; nothing was imported.", [x for x in result if x["errors"]][:50])
        audit.record(db, action="IMPORT_SCHOOLS", entity="organization", entity_id="bulk", user=p.user,
                     after={"created": created, "updated": updated, "file": file.filename}, request=request)
        db.commit()
    return {"committed": commit and ok, "rows": len(result), "create": sum(x["action"] == "create" for x in result),
            "update": sum(x["action"] == "update" for x in result), "errors": sum(x["action"] == "error" for x in result), "lines": result}


# ---------------- user import ----------------
@router.post("/import/users")
async def import_users(request: Request, commit: bool = False, file: UploadFile = File(...), p: Principal = Depends(require("iam:create")),
                       db: Session = Depends(get_db)):
    """CSV columns: full_name, role, org_code and email and/or phone. Each new user gets a temporary password by SMS or
    email and must change it at first sign-in. Existing users get the extra role (segregation-of-duties rules apply)."""
    rows = await _read_csv(file, "users")
    if not rows or not {"full_name", "role", "org_code"} <= set(rows[0]) or not ({"email", "phone"} & set(rows[0])):
        raise AppError(422, "VALIDATION_ERROR", "Columns needed: full_name, role, org_code and email and/or phone.")
    roles = {r.key: r for r in db.scalars(select(Role)).all()}
    staff_only = {"supplier", "farmer_group", "aggregator"}      # suppliers self-register (FR-SUP-01)
    from app.services.admin_access import ADMIN_ROLES
    blocked = ADMIN_ROLES      # administrator accounts are created one by one under Users & roles
    seen: dict[str, set] = {}
    result, created, added = [], 0, 0
    for i, r in enumerate(rows, start=2):
        errs = []
        role = roles.get(r["role"].strip().lower())
        if role is None:
            errs.append(f"Unknown role '{r['role']}'")
        elif role.key in staff_only:
            errs.append("Suppliers register themselves on the website")
        elif role.key in blocked:
            errs.append("Administrator accounts are created under Users & roles, not by import")
        org = db.scalar(select(Organization).where(Organization.code == r["org_code"].upper())) if r.get("org_code") else None
        if r.get("org_code") and org is None:
            errs.append(f"Unknown organisation code '{r['org_code']}'")
        email = (r.get("email") or "").lower() or None
        phone = None
        try:
            phone = normalize_phone(r.get("phone") or None)
        except ValueError as e:
            errs.append(str(e))
        if not email and not phone:
            errs.append("Email or phone is required")
        if email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            errs.append("Invalid email")
        if len(r.get("full_name", "")) < 3:
            errs.append("Full name is required")
        key = email or phone or f"line{i}"
        user = db.scalar(select(User).where((User.email == email) if email else (User.phone == phone))) if (email or phone) else None
        have = set(seen.get(key, set())) | ({ur.role.key for ur in user.roles if ur.is_active} if user else set())
        if have & ADMIN_ROLES:
            errs.append("This is an administrator account; it cannot hold other roles")
        if role:
            for a, b in CONFLICTING_ROLE_PAIRS:
                if (role.key == a and b in have) or (role.key == b and a in have):
                    errs.append(f"{role.key} conflicts with {b if role.key == a else a} (segregation of duties)")
            seen.setdefault(key, set()).add(role.key)
        action = "error" if errs else ("add_role" if user else "create")
        result.append({"line": i, "full_name": r.get("full_name"), "contact": email or notify.mask_phone(phone), "role": r["role"],
                       "org": org.name if org else None, "action": action, "errors": errs})
        if errs or not commit:
            continue
        if user is None:
            temp = new_opaque_token()[:14] + "A1"
            user = User(full_name=r["full_name"], email=email, phone=phone, password_hash=hash_password(temp), status=UserStatus.active,
                        created_by=p.id)
            db.add(user)
            db.flush()
            msg = f"You have been invited to LisheBora as {role.name}. Temporary password: {temp} (change it after signing in)."
            if email:
                notify.email(db, email, "LisheBora invitation", msg, sensitive=True)
            else:
                notify.sms(db, phone, msg, sensitive=True)
            created += 1
        else:
            added += 1
        db.add(UserRole(user_id=user.id, role_id=role.id, org_id=org.id if org else None, created_by=p.id))
        db.flush()
    ok = all(x["action"] != "error" for x in result)
    if commit:
        if not ok:
            db.rollback()
            raise AppError(422, "IMPORT_ERRORS", "Fix the errors in the file first; nothing was imported.", [x for x in result if x["errors"]][:50])
        audit.record(db, action="IMPORT_USERS", entity="user", entity_id="bulk", user=p.user,
                     after={"created": created, "roles_added": added, "file": file.filename}, request=request)
        db.commit()
    return {"committed": commit and ok, "rows": len(result), "create": sum(x["action"] == "create" for x in result),
            "add_role": sum(x["action"] == "add_role" for x in result), "errors": sum(x["action"] == "error" for x in result), "lines": result}


# ---------------- reference price import ----------------
@router.post("/import/prices")
async def import_prices(request: Request, commit: bool = False, file: UploadFile = File(...), p: Principal = Depends(require("md:edit")),
                        db: Session = Depends(get_db)):
    """Columns: code, reference_price (KES per unit). Blank price = leave unchanged. Used for budget estimates and
    price-reasonableness checks; set from county market surveys."""
    from decimal import Decimal, InvalidOperation
    from app.models import Commodity
    rows = await _read_csv(file, "prices")
    if not rows or not {"code", "reference_price"} <= set(rows[0]):
        raise AppError(422, "VALIDATION_ERROR", "Columns needed: code, reference_price.")
    comms = {c.code: c for c in db.scalars(select(Commodity)).all()}
    result, n = [], 0
    for i, r in enumerate(rows, start=2):
        errs, code = [], r["code"].strip().upper()
        c = comms.get(code)
        if c is None:
            errs.append(f"Unknown food code '{r['code']}'")
        price = None
        if r["reference_price"]:
            try:
                price = Decimal(r["reference_price"].replace(",", ""))
                if price < 0 or price > 1_000_000:
                    raise InvalidOperation
            except InvalidOperation:
                errs.append("Price must be a number in KES")
        action = "error" if errs else ("skip" if price is None else "update")
        result.append({"line": i, "code": code, "name": c.name if c else "", "price": str(price) if price is not None else "",
                       "action": action, "errors": errs})
        if commit and action == "update":
            c.reference_price = price
            n += 1
    ok = all(x["action"] != "error" for x in result)
    if commit:
        if not ok:
            db.rollback()
            raise AppError(422, "IMPORT_ERRORS", "Fix the errors in the file first; nothing was imported.", [x for x in result if x["errors"]][:50])
        audit.record(db, action="IMPORT_PRICES", entity="commodity", entity_id="bulk", user=p.user, after={"updated": n, "file": file.filename},
                     request=request)
        db.commit()
    return {"committed": commit and ok, "rows": len(result), "create": 0, "update": sum(x["action"] == "update" for x in result),
            "errors": sum(x["action"] == "error" for x in result), "lines": result}


# ---------------- pilot readiness ----------------
@router.get("/readiness")
def readiness(p: Principal = Depends(require_any("iam:edit", "md:edit")), db: Session = Depends(get_db)):
    """Checks what still stands between this installation and a pilot with real data. Each item: area, title, status
    (ok | todo | warn), detail and a link to where it is fixed."""
    from datetime import date
    from app.core.security import verify_password
    from app.models import (AcademicTerm, Commodity, ContentStatus, Menu, MenuStatus, Page, ProcurementEvent, PurchaseOrder,
                            SiteBlock, Supplier)
    items = []

    def add(area, title, ok, detail, link, warn=False):
        items.append({"area": area, "title": title, "status": "ok" if ok else ("warn" if warn else "todo"), "detail": detail, "link": link})

    # security & environment
    admin = db.scalar(select(User).where(User.email == settings.seed_admin_email.lower()))
    add("Security", "Super Administrator password changed", not (admin and verify_password("ChangeMe!2026", admin.password_hash)),
        "The first Super Administrator still uses the default password. Change it under My profile & security." if admin and verify_password("ChangeMe!2026", admin.password_hash)
        else "The default password is no longer in use.", "/app/profile")
    add("Security", "Sign-in codes are not shown on screen", not settings.debug,
        "DEBUG is on, so SMS codes are shown on screen. Fine on a training laptop; turn it off (DEBUG=false) for the pilot server." if settings.debug
        else "Codes are only sent by SMS or email.", "/app/system", warn=True)
    add("Security", "HTTPS-only cookies", settings.cookie_secure, "Set COOKIE_SECURE=true on the pilot server (behind HTTPS)." if not settings.cookie_secure
        else "On.", "/app/system", warn=True)
    add("Messaging", "SMS gateway connected", settings.notify_backend != "console",
        "SMS are only written to the server log. Set NOTIFY_BACKEND=africastalking with the account details." if settings.notify_backend == "console"
        else f"Using {settings.notify_backend}.", "/app/system")
    add("Messaging", "Email server connected", bool(settings.smtp_host), "Email is only written to the server log. Set the SMTP_* settings." if not settings.smtp_host
        else "SMTP configured.", "/app/system")
    add("Infrastructure", "Production database", db.bind.dialect.name == "postgresql",
        "SQLite is fine for a single training laptop. Use PostgreSQL for the pilot server (see docs/DEPLOYMENT.md)." if db.bind.dialect.name != "postgresql"
        else "PostgreSQL.", "/app/system", warn=True)
    jobs = db.scalar(select(SystemSetting).where(SystemSetting.key == "jobs.last_run"))
    add("Infrastructure", "Background worker running", bool(jobs and jobs.value),
        "No background job has run yet: SMS/email are not sent and reminders do not go out. Start scripts\\start-worker.bat (or the worker service)."
        if not (jobs and jobs.value) else "Jobs have run.", "/app/system")
    # demo data
    demo_users = db.scalar(select(func.count()).select_from(User).where(User.email.like("%@demo.lishebora"))) or 0
    demo_po = db.scalar(select(func.count()).select_from(PurchaseOrder).where(PurchaseOrder.reference.like("PO-DEMO-%"))) or 0
    demo_sup = db.scalar(select(func.count()).select_from(Organization).where(Organization.code.like("SUP-DEMO-%"))) or 0
    demo_ev = db.scalar(select(func.count()).select_from(ProcurementEvent).where(ProcurementEvent.title.like("%(demo)%"))) or 0
    clean = not (demo_users or demo_po or demo_sup or demo_ev)
    add("Data", "No demonstration data", clean,
        f"This database contains demo data ({demo_users} demo accounts, {demo_sup} demo suppliers, {demo_po} demo orders, {demo_ev} demo notices). "
        "Create a clean pilot database with scripts\\setup-pilot.bat." if not clean else "Clean.", "/app/users")
    # geography & schools
    counties = db.scalar(select(func.count()).select_from(Organization).where(Organization.type == OrgType.county, Organization.is_active)) or 0
    schools = db.scalars(select(Organization).where(Organization.type == OrgType.school, Organization.is_active)).all()
    add("Data", "Pilot counties and schools", counties > 0 and len(schools) > 0, f"{counties} counties, {len(schools)} schools.", "/app/master-data")
    no_enrol = [s for s in schools if not (s.meta or {}).get("enrolment")]
    add("Data", "Enrolment for every school", not no_enrol, f"{len(no_enrol)} of {len(schools)} schools have no enrolment. Fill the Schools sheet of the pilot workbook and import it."
        if no_enrol else "All schools have enrolment.", "/app/system")
    no_gps = [s for s in schools if not (s.meta or {}).get("gps")]
    add("Data", "GPS location for every school", not no_gps, f"{len(no_gps)} schools have no GPS: they are missing from the map and route planning. "
        "Add lat/lng in the workbook, or school staff capture it on site under School locations." if no_gps else "All mapped.", "/app/locations", warn=True)
    no_nemis = [s for s in schools if not (s.meta or {}).get("nemis_code")]
    add("Data", "NEMIS code for every school", not no_nemis, f"{len(no_nemis)} schools have no NEMIS code (optional, helps matching with Ministry data)."
        if no_nemis else "All recorded.", "/app/master-data", warn=True)
    comms = db.scalars(select(Commodity).where(Commodity.is_active)).all()
    no_price = [c for c in comms if c.reference_price is None]
    add("Data", "Reference prices for foods", not no_price, f"{len(no_price)} of {len(comms)} foods have no reference price, so budget estimates are incomplete. "
        "Fill the Prices sheet and import it." if no_price else "All foods priced.", "/app/system", warn=True)
    # people
    def holders(role, org_ids=None):
        q = select(func.count(func.distinct(UserRole.user_id))).join(Role, Role.id == UserRole.role_id).join(User, User.id == UserRole.user_id).where(
            Role.key == role, UserRole.is_active, User.status == UserStatus.active, ~User.email.like("%@demo.lishebora") | User.email.is_(None))
        if org_ids is not None:
            q = q.where(UserRole.org_id.in_(org_ids))
        return db.scalar(q) or 0
    admins = holders("system_admin")
    add("People", "At least one Administrator (besides the Super Administrator)", admins > 0, f"{admins} Administrator accounts.", "/app/users")
    county_rows = db.scalars(select(Organization).where(Organization.type == OrgType.county, Organization.is_active)).all()
    for role, label in (("county_admin", "County Administrator"), ("county_procurement_officer", "County Procurement Officer"),
                        ("county_finance_officer", "County Finance Officer"), ("nutrition_officer", "Nutrition Officer")):
        missing = [c.name for c in county_rows if holders(role, [c.id]) == 0]
        add("People", f"{label} in every pilot county", not missing, ("Missing in: " + ", ".join(missing)) if missing else "Assigned.", "/app/users")
    add("People", "Approving Officer", holders("approving_officer") > 0, "Needed to publish notices and approve awards.", "/app/users")
    add("People", "Evaluation committee (at least 3 members)", holders("evaluation_committee_member") >= 3,
        f"{holders('evaluation_committee_member')} members.", "/app/users", warn=True)
    sch_ids = [s.id for s in schools]
    with_admin = {r for r in db.scalars(select(UserRole.org_id).join(Role, Role.id == UserRole.role_id).join(User, User.id == UserRole.user_id).where(
        Role.key == "school_admin", UserRole.is_active, UserRole.org_id.in_(sch_ids or [None]), ~User.email.like("%@demo.lishebora")))}
    add("People", "A School Administrator for every school", len(with_admin) >= len(schools),
        f"{len(schools) - len(with_admin)} of {len(schools)} schools have no School Administrator. Import the school contacts "
        "(docs/data/school_contacts_for_import.csv) once the heads agree." if len(with_admin) < len(schools) else "All schools covered.", "/app/system")
    for role, label in (("quality_inspector", "Quality Inspector"), ("warehouse_officer", "Warehouse Officer"), ("logistics_officer", "Logistics Officer"),
                        ("finance_officer", "Finance Officer (verifies invoices, records payments)")):
        add("People", label, holders(role) > 0, "Needed for deliveries and payments.", "/app/users", warn=role != "finance_officer")
    # planning
    this_year = date.today().year
    terms = db.scalars(select(AcademicTerm).where(AcademicTerm.ends_on >= date.today())).all()
    add("Planning", "School terms for the pilot period", len(terms) > 0, f"{len(terms)} current or future terms. Check the dates against the Ministry calendar under Menus & terms."
        if terms else f"No current or future term. Add the {this_year}/{this_year + 1} terms under Menus & terms.", "/app/menus")
    menus = db.scalar(select(func.count()).select_from(Menu).where(Menu.status == MenuStatus.approved, ~Menu.name.like("%(demo)%"))) or 0
    add("Planning", "An approved (non-demo) menu", menus > 0, "Nutrition officers create and approve the pilot menu under Menus & terms." if not menus
        else f"{menus} approved menus.", "/app/menus")
    # website
    placeholder_pages = [pg.title for pg in db.scalars(select(Page)).all() if "*[" in (pg.body or "") or "to be supplied" in (pg.body or "") or "Draft for the pilot" in (pg.body or "")]
    add("Website", "Pages without placeholder text", not placeholder_pages, ("Still has placeholders: " + ", ".join(placeholder_pages)) if placeholder_pages
        else "All pages complete.", "/app/website")
    site = db.scalar(select(SiteBlock).where(SiteBlock.key == "site"))
    sd = (site.published_data or {}) if site else {}
    add("Website", "Helpdesk contact details", bool(sd.get("phone") or sd.get("email")), "Add the helpdesk phone, SMS short code and email under Website content → Homepage → Contact details."
        if not (sd.get("phone") or sd.get("email")) else "Published.", "/app/website")
    stats = db.scalar(select(SiteBlock).where(SiteBlock.key == "stats"))
    st = (stats.published_data or {}) if stats else {}
    typed_placeholder = st.get("show") and any("[x]" in str(i.get("value", "")) for i in st.get("items", []))
    add("Website", "Programme statistics are real figures", not typed_placeholder, "The homepage shows '[x]' placeholders. Enter signed-off MEAL figures or hide the block."
        if typed_placeholder else "OK.", "/app/website")
    suppliers = db.scalar(select(func.count()).select_from(Supplier).where(~Supplier.email.like("%@demo.lishebora") | Supplier.email.is_(None))) or 0
    add("Suppliers", "Suppliers have started registering", suppliers > 0, f"{suppliers} supplier registrations (excluding demo).", "/app/suppliers", warn=True)
    order = {"todo": 0, "warn": 1, "ok": 2}
    return {"summary": {k: sum(i["status"] == k for i in items) for k in ("todo", "warn", "ok")},
            "items": sorted(items, key=lambda i: (order[i["status"]], i["area"]))}


# ---------------- pilot data workbook ----------------
@router.get("/workbook")
def pilot_workbook(p: Principal = Depends(require_any("iam:edit", "md:edit")), db: Session = Depends(get_db)):
    """An Excel workbook pre-filled with what the platform already holds, for the programme team to complete and import:
    Schools (enrolment, GPS, NEMIS), Staff accounts, Prices, plus reference lists of roles and organisation codes."""
    from fastapi.responses import Response
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from app.models import Commodity
    from app.seed.matrix import ROLES
    from app.seed.reference import CATEGORY_LABELS
    from app.services.admin_access import ADMIN_ROLES
    wb = Workbook()
    head = PatternFill("solid", fgColor="2F4520"); fill = PatternFill("solid", fgColor="FEF6DD")
    def sheet(title, note, cols, rows, widths, editable=()):
        ws = wb.create_sheet(title)
        ws.append([note]); ws["A1"].font = Font(italic=True, color="4B5842")
        ws.append(cols)
        for c in ws[2]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = head
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = w
        for col in editable:   # cells the team should fill are shaded
            idx = cols.index(col) + 1
            for row in ws.iter_rows(min_row=3, max_row=max(3, ws.max_row + 200), min_col=idx, max_col=idx):
                row[0].fill = fill
        ws.freeze_panes = "A3"
        return ws
    orgs = db.scalars(select(Organization).where(Organization.is_active)).all()
    by_id = {o.id: o for o in orgs}
    def county_of(o):
        for _ in range(6):
            if o is None or o.type == OrgType.county:
                return o
            o = by_id.get(o.parent_id)
    schools = sorted((o for o in orgs if o.type == OrgType.school), key=lambda o: o.code)
    rows = []
    for s in schools:
        parent = by_id.get(s.parent_id); sub = parent if parent and parent.type == OrgType.sub_county else by_id.get(parent.parent_id) if parent else None
        cl = parent.name if parent and parent.type == OrgType.school_cluster else ""
        m = s.meta or {}; g = m.get("gps") or {}
        c = county_of(s)
        rows.append([c.name if c else "", sub.name if sub else "", cl, s.code, s.name, m.get("enrolment"), g.get("lat"), g.get("lng"), m.get("nemis_code", "")])
    wb.remove(wb.active)
    ins = wb.create_sheet("Instructions")
    for line in ["LisheBora pilot data workbook",
                 "1. Fill the shaded cells in the Schools, Staff and Prices sheets. Do not change the column headings in row 2.",
                 "2. Schools: enrolment (number of learners), GPS in decimal degrees (e.g. -2.3551 and 37.9012), and the NEMIS code if known.",
                 "   Keep school_code as it is: it links the school to its records. To add a school, add a row with a new code (e.g. MAKUENI-SCH017).",
                 "3. Staff: one row per person and role. org_code is a code from the 'Codes' sheet (a school, sub-county or county).",
                 "   Roles are listed in the 'Roles' sheet. Give either an email or a mobile number; the person gets a temporary password by SMS or email.",
                 "   Suppliers do not go here: they register themselves on the website. Administrator accounts are created one by one under Users & roles.",
                 "4. Prices: reference price in KES per unit (kg, litre or tray) from the county market survey. Leave blank if not known.",
                 "5. Upload the whole workbook under System & onboarding: Import schools, Import staff accounts and Import reference prices.",
                 "   Each import reads its own sheet, shows a preview, and saves nothing until the file has no errors.",
                 "Personal data (names, phone numbers) is covered by the Data Protection Act, 2019: share this file only with the programme team."]:
        ins.append([line])
    ins["A1"].font = Font(bold=True, size=14, color="2F4520"); ins.column_dimensions["A"].width = 130
    sheet("Schools", "Schools: fill enrolment, lat, lng and nemis_code (shaded).",
          ["county", "sub_county", "cluster", "school_code", "school_name", "enrolment", "lat", "lng", "nemis_code"], rows,
          [18, 18, 14, 18, 34, 12, 12, 12, 14], editable=("enrolment", "lat", "lng", "nemis_code"))
    roles = [(k, v[0]) for k, v in ROLES.items() if k not in ADMIN_ROLES and k not in ("supplier", "farmer_group", "aggregator")]
    st = sheet("Staff", "Staff accounts: one row per person and role.", ["full_name", "email", "phone", "role", "org_code"], [],
               [30, 30, 16, 30, 22], editable=("full_name", "email", "phone", "role", "org_code"))
    dv = DataValidation(type="list", formula1=f"=Roles!$A$3:$A${len(roles) + 2}", allow_blank=True); st.add_data_validation(dv); dv.add("D3:D1000")
    comms = db.scalars(select(Commodity).where(Commodity.is_active).order_by(Commodity.category, Commodity.name)).all()
    sheet("Prices", "Reference prices in KES per unit (shaded).", ["code", "name", "food_category", "unit", "reference_price"],
          [[c.code, c.name, CATEGORY_LABELS.get(c.category, c.category), c.unit, float(c.reference_price) if c.reference_price is not None else None] for c in comms],
          [16, 22, 26, 8, 16], editable=("reference_price",))
    sheet("Roles", "Role codes for the Staff sheet.", ["role", "name", "typically"],
          [[k, n, ""] for k, n in roles], [30, 32, 30])
    code_rows = [[o.code, o.name, o.type.value.replace("_", " "), (county_of(o).name if county_of(o) else "")] for o in
                 sorted(orgs, key=lambda o: (o.type.value, o.code)) if o.type in (OrgType.county, OrgType.sub_county, OrgType.school_cluster, OrgType.school, OrgType.warehouse)]
    sheet("Codes", "Organisation codes for the Staff sheet (org_code).", ["org_code", "name", "type", "county"], code_rows, [22, 34, 14, 18])
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(vertical="top")
    buf = io.BytesIO(); wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="LisheBora_Pilot_Data_Workbook.xlsx"'})
