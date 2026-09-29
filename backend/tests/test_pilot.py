"""Pilot readiness: readiness checks, Excel imports (schools with enrolment/NEMIS, prices), site contact details."""
import io

from openpyxl import Workbook
from sqlalchemy import select

from app.core.database import SessionLocal
from app.models import Commodity, Organization
from tests.conftest import login

A = "/api/v1"
ADMIN = ("administrator@demo.lishebora", "Demo!2026pass")


def _xlsx(sheets):
    wb = Workbook(); wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def test_readiness_flags_demo_and_gaps(client):
    ah = login(client, *ADMIN)
    r = client.get(f"{A}/system/readiness", headers=ah).json()
    items = {i["title"]: i for i in r["items"]}
    assert items["No demonstration data"]["status"] == "todo"
    assert items["Enrolment for every school"]["status"] == "todo"
    assert items["Helpdesk contact details"]["status"] == "todo"
    assert r["summary"]["todo"] >= 3 and all(i["link"].startswith("/") for i in r["items"])
    assert client.get(f"{A}/system/readiness", headers=login(client, "procurement@demo.lishebora", "Demo!2026pass")).status_code == 403


def test_excel_imports(client):
    ah = login(client, *ADMIN)
    book = _xlsx({"Instructions": [["Fill the sheets"]],
                  "Schools": [["Pilot schools"], ["county", "sub_county", "cluster", "school_code", "school_name", "enrolment", "lat", "lng", "nemis_code"],
                              ["Makueni County", "Kibwezi West", None, "MAKUENI-SCH003", "Mukononi", 431, -2.31, 37.95, "ab12345"]],
                  "Prices": [["code", "name", "reference_price"], ["BEANS", "Beans", "135.5"], ["SALT", "Salt", None], ["NOPE", "x", 3]]})
    files = lambda: {"file": ("pilot.xlsx", book, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = client.post(f"{A}/system/import/schools?commit=true", headers=ah, files=files()).json()
    assert r["committed"] and r["update"] == 1, r
    with SessionLocal() as db:
        s = db.scalar(select(Organization).where(Organization.code == "MAKUENI-SCH003"))
        assert s.meta["enrolment"] == 431 and s.meta["nemis_code"] == "AB12345" and s.meta["gps"]["lat"] == -2.31
    r = client.post(f"{A}/system/import/prices", headers=ah, files=files()).json()
    assert r["errors"] == 1 and [l["action"] for l in r["lines"]] == ["update", "skip", "error"]
    ok = _xlsx({"Prices": [["code", "reference_price"], ["BEANS", "135.5"], ["SALT", ""]]})
    r = client.post(f"{A}/system/import/prices?commit=true", headers=ah, files={"file": ("p.xlsx", ok, "application/octet-stream")}).json()
    assert r["committed"] and r["update"] == 1
    with SessionLocal() as db:
        assert str(db.scalar(select(Commodity).where(Commodity.code == "BEANS")).reference_price) in ("135.5", "135.50")
    reach = {c["code"]: c for c in client.get(f"{A}/public/reach").json()["counties"]}
    assert reach["MAKUENI"]["pupils"] >= 431


def test_site_contact_details(client):
    ah = login(client, *ADMIN)
    b = client.put(f"{A}/cms/blocks/site", headers=ah, json={"data": {"phone": " 0700 123 456 ", "email": "help@example.org", "evil": "<x>"}}).json()
    assert b["data"]["phone"] == "0700 123 456" and "evil" not in b["data"]
    client.post(f"{A}/cms/blocks/site/workflow", headers=ah, json={"action": "publish"})
    assert client.get(f"{A}/public/site").json()["email"] == "help@example.org"


def test_workbook_round_trip(client):
    from openpyxl import load_workbook
    ah = login(client, *ADMIN)
    r = client.get(f"{A}/system/workbook", headers=ah)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxml")
    wb = load_workbook(io.BytesIO(r.content))
    assert {"Instructions", "Schools", "Staff", "Prices", "Roles", "Codes"} <= set(wb.sheetnames)
    assert wb["Schools"].max_row >= 34 and wb["Schools"]["D2"].value == "school_code"
    # the untouched workbook imports cleanly (schools update in place, no errors)
    res = client.post(f"{A}/system/import/schools", headers=ah, files={"file": ("w.xlsx", r.content, "application/octet-stream")}).json()
    assert res["errors"] == 0 and res["create"] == 0 and res["update"] >= 32, res



def test_school_locations(client):
    """School staff set their own school's GPS (phone, typed or map); coordinates outside Kenya are refused;
    the list shows who captured it; only administrators can clear."""
    from tests.conftest import org_id
    from tests.test_phase4 import A, h
    sa, adm = h(client, "school.admin"), login(client, *ADMIN)
    sid = str(org_id("MAKUENI-SCH001"))
    r = client.get(f"{A}/fulfilment/gps-locations", headers=sa).json()
    assert [x["id"] for x in r["items"]] == [sid]                      # a school admin sees only their school
    assert client.put(f"{A}/fulfilment/locations/{sid}/gps", headers=sa, json={"lat": 37.6, "lng": -1.8}).status_code == 422   # swapped
    ok = client.put(f"{A}/fulfilment/locations/{sid}/gps", headers=sa,
                    json={"lat": -1.803512, "lng": 37.620145, "accuracy_m": 8, "source": "device", "note": "kitchen gate"})
    assert ok.status_code == 200 and ok.json()["gps"]["source"] == "device" and ok.json()["gps"]["by"]
    item = next(x for x in client.get(f"{A}/fulfilment/gps-locations?q=MAKUENI-SCH001", headers=adm).json()["items"])
    assert item["gps"]["note"] == "kitchen gate" and item["county"] == "Makueni County"
    assert client.delete(f"{A}/fulfilment/locations/{sid}/gps", headers=sa).status_code == 403
    assert client.delete(f"{A}/fulfilment/locations/{sid}/gps", headers=adm).status_code == 200
    assert client.get(f"{A}/fulfilment/gps-locations?missing=true&q=MAKUENI-SCH001", headers=adm).json()["total"] == 1
