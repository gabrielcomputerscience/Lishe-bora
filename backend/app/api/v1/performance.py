"""Supplier performance scorecards and the MEAL dashboard with CSV exports (SRS §3.12–3.13, FR-MEAL-01…04)."""
import csv
import io
import uuid
from datetime import date
from decimal import Decimal

from fastapi import APIRouter, Depends, Request
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.core.database import get_db
from app.core.deps import Principal, require, require_any, scope_any, scope_for
from app.core.errors import AppError
from app.core.timeutil import as_utc
from app.models import (Complaint, Dispatch, DispatchStatus, ExceptionCase, Intake, Invoice, Organization, OrgType, Payment,
                        ProofOfDelivery, PurchaseOrder, Supplier)
from app.api.v1.finance import sweep_late_payments
from app.services import audit, performance as perf
from app.services.stock import county_of

router = APIRouter(tags=["performance & MEAL"])


def _own_supplier(db, p: Principal) -> Supplier | None:
    orgs = {ur.org_id for ur in p.user.roles if ur.is_active and ur.org_id}
    return db.scalar(select(Supplier).where(Supplier.organization_id.in_(orgs))) if orgs else None


def _suppliers_in_scope(db, p: Principal) -> list[Supplier]:
    own = _own_supplier(db, p)
    if own:
        return [own]
    allowed = scope_any(db, p, "sup:view", "cmp:view", "meal:view")
    q = select(Supplier)
    if allowed is not None:
        q = q.where(Supplier.county_id.in_(allowed))
    return list(db.scalars(q.order_by(Supplier.legal_name)).all())


@router.get("/performance/suppliers")
def supplier_scores(p: Principal = Depends(require_any("sup:view", "cmp:view", "meal:view")), db: Session = Depends(get_db)):
    rows = [perf.scorecard(db, s) for s in _suppliers_in_scope(db, p)]
    return sorted(rows, key=lambda r: (r["score"] is None, -(r["score"] or 0), r["supplier"]))


@router.get("/performance/suppliers/{sid}")
def supplier_score(sid: uuid.UUID, p: Principal = Depends(require_any("sup:view", "cmp:view", "meal:view")), db: Session = Depends(get_db)):
    s = next((x for x in _suppliers_in_scope(db, p) if x.id == sid), None)
    if s is None:
        raise AppError(404, "NOT_FOUND", "Supplier not found in your area.")
    card = perf.scorecard(db, s)
    po_ids = select(PurchaseOrder.id).where(PurchaseOrder.supplier_id == s.id)
    dispatches = {d.id: d for d in db.scalars(select(Dispatch).where(Dispatch.po_id.in_(po_ids),
                                                                     Dispatch.status != DispatchStatus.cancelled)).all()}
    pods = db.scalars(select(ProofOfDelivery).where(ProofOfDelivery.dispatch_id.in_(list(dispatches) or [uuid.uuid4()]))
                      .order_by(ProofOfDelivery.received_at.desc()).limit(30)).all()
    schools = {o.id: o.name for o in db.scalars(select(Organization).where(Organization.type == OrgType.school)).all()}
    card["recent_deliveries"] = [{"dispatch": dispatches[x.dispatch_id].reference, "school": schools.get(x.school_id, ""),
                                  "planned": dispatches[x.dispatch_id].planned_date, "received_at": x.received_at, "status": x.status,
                                  "on_time": perf._on_time(x, dispatches[x.dispatch_id])} for x in pods]
    card["recent_complaints"] = [{"reference": c.reference, "category": c.category, "subject": c.subject, "status": c.status.value,
                                  "created_at": c.created_at} for c in db.scalars(select(Complaint).where(
                                      Complaint.supplier_id == s.id).order_by(Complaint.created_at.desc()).limit(10))
                                 if not (_own_supplier(db, p) and c.category in ("fraud", "safeguarding", "staff_conduct"))]
    return card


# ---------------- MEAL ----------------
def _counties(db, p: Principal, code: str, county_id: uuid.UUID | None) -> set[uuid.UUID] | None:
    allowed = scope_for(db, p, code)
    if county_id:
        if allowed is not None and county_id not in allowed:
            raise AppError(403, "OUT_OF_SCOPE", "This county is outside your assigned area.")
        return {county_id}
    if allowed is None:
        return None
    return {c for c in (county_of(db, o) for o in allowed) if c} if allowed else set()


@router.get("/meal/dashboard")
def meal_dashboard(county_id: uuid.UUID | None = None, start: date | None = None, end: date | None = None,
                   p: Principal = Depends(require("meal:view")), db: Session = Depends(get_db)):
    sweep_late_payments(db)
    counties = _counties(db, p, "meal:view", county_id)
    out = perf.meal_dashboard(db, counties, start, end)
    names = db.scalars(select(Organization).where(Organization.type == OrgType.county)).all()
    mine = _counties(db, p, "meal:view", None)     # the filter list shows every county the user may pick
    out["counties"] = [{"id": c.id, "name": c.name} for c in names if mine is None or c.id in mine]
    out["scope"] = "All counties" if counties is None else ", ".join(c.name for c in names if c.id in counties) or "your area"
    return out


DATASETS = {
    "deliveries": "School deliveries (proof of delivery)",
    "invoices": "Invoices and payments",
    "suppliers": "Supplier scorecards",
    "producers": "Producer participation by hub (aggregated, no personal data)",
    "exceptions": "Exception cases",
    "complaints": "Complaints (no contact details)",
}


@router.get("/meal/exports")
def exports(p: Principal = Depends(require("meal:export"))):
    return [{"key": k, "label": v} for k, v in DATASETS.items()]


def _rows(db: Session, p: Principal, key: str, counties: set | None) -> tuple[list[str], list[list]]:
    inc = (lambda c: counties is None or c in counties)
    if key == "deliveries":
        schools = {o.id: o.name for o in db.scalars(select(Organization).where(Organization.type == OrgType.school)).all()}
        disp = {d.id: d for d in db.scalars(select(Dispatch).options(selectinload(Dispatch.lines))).all() if inc(d.county_id)}
        out = []
        for x in db.scalars(select(ProofOfDelivery).order_by(ProofOfDelivery.received_at)).all():
            d = disp.get(x.dispatch_id)
            if not d:
                continue
            for ln in d.lines:
                if ln.school_id == x.school_id:
                    out.append([d.reference, schools.get(x.school_id, ""), ln.commodity_code, ln.unit, ln.quantity, ln.accepted_qty,
                                ln.rejected_qty, ln.rejection_reason, d.planned_date, as_utc(x.received_at).date(),
                                "yes" if perf._on_time(x, d) else "no", x.status, "yes" if x.captured_offline else "no"])
        return (["dispatch", "school", "commodity", "unit", "sent", "accepted", "rejected", "rejection_reason", "planned_date",
                 "received_date", "on_time", "pod_status", "captured_offline"], out)
    if key == "invoices":
        sups = {s.id: s.legal_name for s in db.scalars(select(Supplier)).all()}
        pays: dict = {}
        for pm in db.scalars(select(Payment)).all():
            pays.setdefault(pm.invoice_id, []).append(pm)
        out = []
        for i in db.scalars(select(Invoice).options(selectinload(Invoice.po)).order_by(Invoice.submitted_at)).all():
            if not inc(i.county_id):
                continue
            pms = pays.get(i.id, [])
            out.append([i.reference, i.supplier_invoice_no, sups.get(i.supplier_id, ""), i.po.reference if i.po else "", i.invoice_date,
                        i.total, i.paid_amount, i.status.value, (i.match or {}).get("result", ""), as_utc(i.submitted_at).date(),
                        as_utc(i.approved_at).date() if i.approved_at else "", as_utc(i.paid_at).date() if i.paid_at else "",
                        (as_utc(i.paid_at) - as_utc(i.submitted_at)).days if i.paid_at else "", "; ".join(f"{x.reference} {x.method} {x.transaction_ref}" for x in pms)])
        return (["invoice", "supplier_invoice_no", "supplier", "po", "invoice_date", "total", "paid", "status", "three_way_match",
                 "submitted", "approved", "paid_on", "days_to_pay", "payments"], out)
    if key == "suppliers":
        out = []
        for s in db.scalars(select(Supplier)).all():
            if not inc(s.county_id):
                continue
            c = perf.scorecard(db, s)
            r = c["rates"]
            out.append([c["supplier"], c["supplier_type"], c["inclusion"] or "", c["orders"], c["deliveries"], r["on_time"], r["acceptance"],
                        r["fill"], r["inspection"], c["complaints"], c["avg_days_to_pay"], c["score"], c["band"] or ""])
        return (["supplier", "type", "verified_inclusion", "orders", "deliveries", "on_time_pct", "acceptance_pct", "fill_pct",
                 "inspection_pass_pct", "complaints", "avg_days_to_pay", "score", "band"], out)
    if key == "producers":
        hubs = {o.id: o for o in db.scalars(select(Organization).where(Organization.type == OrgType.aggregation_centre)).all()}
        agg: dict = {}
        for i in db.scalars(select(Intake)).all():
            h = hubs.get(i.location_id)
            if h is None or not inc(county_of(db, h.id)):
                continue
            a = agg.setdefault((h.name, i.commodity_code), {"producers": set(), "women": set(), "youth": set(), "qty": Decimal(0), "intakes": 0})
            k = (i.producer_name.strip().lower(), i.producer_phone)
            a["producers"].add(k)
            if i.producer_gender == "female":
                a["women"].add(k)
            if i.producer_youth:
                a["youth"].add(k)
            a["qty"] += Decimal(i.quantity)
            a["intakes"] += 1
        return (["hub", "commodity", "producers", "women_producers", "youth_producers", "intakes", "quantity_kg"],
                [[h, c, len(a["producers"]), len(a["women"]), len(a["youth"]), a["intakes"], a["qty"]] for (h, c), a in sorted(agg.items())])
    if key == "exceptions":
        return (["reference", "category", "severity", "entity", "title", "status", "owner_role", "raised", "due", "resolved"],
                [[c.reference, c.category, c.severity, c.entity_ref, c.title, c.status.value, c.owner_role, as_utc(c.created_at).date(),
                  as_utc(c.due_at).date() if c.due_at else "", as_utc(c.resolved_at).date() if c.resolved_at else ""]
                 for c in db.scalars(select(ExceptionCase).order_by(ExceptionCase.created_at)).all() if counties is None or c.org_id in counties])
    if key == "complaints":
        return (["reference", "channel", "category", "priority", "subject", "status", "raised", "due", "resolved", "satisfaction"],
                [[c.reference, c.channel, c.category, c.priority, c.subject, c.status.value, as_utc(c.created_at).date(),
                  as_utc(c.due_at).date() if c.due_at else "", as_utc(c.resolved_at).date() if c.resolved_at else "", c.satisfaction or ""]
                 for c in db.scalars(select(Complaint).order_by(Complaint.created_at)).all() if counties is None or c.org_id in counties])
    raise AppError(404, "NOT_FOUND", "Unknown dataset.")


@router.get("/meal/exports/{key}.xlsx")
def export_xlsx(key: str, request: Request, county_id: uuid.UUID | None = None, p: Principal = Depends(require("meal:export")),
                db: Session = Depends(get_db)):
    """Same datasets as the CSV exports, as a formatted Excel workbook (FR-MEAL-03)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    if key not in DATASETS:
        raise AppError(404, "NOT_FOUND", "Unknown dataset.")
    counties = _counties(db, p, "meal:export", county_id)
    head, rows = _rows(db, p, key, counties)
    wb = Workbook()
    ws = wb.active
    ws.title = key[:31]
    ws.append([DATASETS[key]])
    ws["A1"].font = Font(bold=True, size=13, color="3D5A27")
    ws.append([f"LisheBora · exported {date.today():%d %b %Y} · generated from platform records"])
    ws.append([])
    ws.append([h.replace("_", " ").capitalize() for h in head])
    for c in ws[4]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4F7A2E")
    for r in rows:
        ws.append([float(v) if isinstance(v, Decimal) else (str(v) if isinstance(v, uuid.UUID) else v) for v in r])
    for i, h in enumerate(head, start=1):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = max(12, min(40, len(h) + 4))
    ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    audit.record(db, action="EXPORT", entity="dataset", entity_id=key, user=p.user, after={"rows": len(rows), "format": "xlsx"}, request=request)
    db.commit()
    return StreamingResponse(iter([buf.getvalue()]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="lishebora_{key}_{date.today():%Y%m%d}.xlsx"'})


@router.get("/meal/exports/{key}.csv")
def export_csv(key: str, request: Request, county_id: uuid.UUID | None = None, p: Principal = Depends(require("meal:export")),
               db: Session = Depends(get_db)):
    if key not in DATASETS:
        raise AppError(404, "NOT_FOUND", "Unknown dataset.")
    counties = _counties(db, p, "meal:export", county_id)
    head, rows = _rows(db, p, key, counties)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(head)
    w.writerows(rows)
    audit.record(db, action="EXPORT", entity="dataset", entity_id=key, user=p.user, after={"rows": len(rows), "county": county_id}, request=request)
    db.commit()
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
                             headers={"Content-Disposition": f'attachment; filename="lishebora_{key}_{date.today():%Y%m%d}.csv"'})
